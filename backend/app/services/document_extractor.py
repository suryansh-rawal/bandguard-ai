import io
from typing import Tuple
from pypdf import PdfReader
from docx import Document as DocxDocument

class DocumentExtractor:
    """
    K12: Document text extraction utility supporting DOCX, PDF, TXT, and Markdown files.
    """

    @staticmethod
    def extract_text(filename: str, content_bytes: bytes) -> Tuple[str, str]:
        """
        Extract plain text from file bytes.
        Returns: (extracted_text, detected_type)
        """
        ext = filename.split(".")[-1].lower() if "." in filename else "txt"

        if ext in ["txt", "md", "csv", "log"]:
            try:
                text = content_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = content_bytes.decode("latin-1", errors="replace")
            return text.strip(), ext

        elif ext == "docx":
            try:
                doc = DocxDocument(io.BytesIO(content_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also extract table text if present
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                return "\n\n".join(paragraphs).strip(), "docx"
            except Exception as e:
                raise ValueError(f"Failed to parse DOCX file: {e}")

        elif ext == "pdf":
            try:
                reader = PdfReader(io.BytesIO(content_bytes))
                pages = []
                for idx, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        pages.append(page_text.strip())
                return "\n\n".join(pages).strip(), "pdf"
            except Exception as e:
                raise ValueError(f"Failed to parse PDF file: {e}")

        else:
            # Fallback to UTF-8 text decode
            try:
                text = content_bytes.decode("utf-8")
                return text.strip(), "unknown"
            except Exception:
                raise ValueError(f"Unsupported file format: {ext}")

document_extractor = DocumentExtractor()
